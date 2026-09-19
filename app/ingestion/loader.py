"""Load supported enterprise document formats into a common RAG record shape.

Each loader returns a list of dictionaries containing ``text`` and ``metadata``
keys. This normalized structure lets the downstream chunker and vector store
handle text files, PDFs, Office documents, images, and tabular files uniformly.
PDF pages and images fall back to OCR when their native text is unavailable.
"""

from pathlib import Path

import pandas as pd
import pytesseract
from PIL import Image
from docx import Document
import pymupdf


def load_txt(file_path: Path):
    """Read a UTF-8 text file and return it as one normalized document record.

    Args:
        file_path: Path to the ``.txt`` source.

    Returns:
        A one-item list with source, type, page, and section metadata.
    """
    text = file_path.read_text(encoding="utf-8")

    return [
        {
            "text": text,
            "metadata": {
                "source": file_path.name,
                "file_type": "txt",
                "page":None,
                "section":None
            },
        }
    ]


def load_pdf(file_path: Path):
    """Extract one normalized record per non-empty PDF page.

    Native page text is preferred. For scanned pages that contain no extractable
    text, the page is rasterized at 2× scale and passed through Tesseract OCR.

    Args:
        file_path: Path to a valid PDF source.

    Returns:
        Normalized records with one-based page numbers.

    Raises:
        ValueError: If a file with a ``.pdf`` extension lacks the PDF signature.
    """
    documents = []

    # A file named .pdf must contain the PDF signature, not merely text with a
    # PDF extension. Checking it here gives a clear remediation path.
    with file_path.open("rb") as file:
        pdf_signature = file.read(5)

    if pdf_signature != b"%PDF-":
        raise ValueError(
            "File has a .pdf extension but is not a valid PDF. "
            "Convert it to PDF or rename it with a .txt extension."
        )

    pdf = pymupdf.open(file_path)

    for page_number, page in enumerate(pdf, start=1):
        text = page.get_text().strip()

        # If the PDF is scanned, use OCR
        if not text:
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
            image = Image.frombytes(
                "RGB",
                [pixmap.width, pixmap.height],
                pixmap.samples,
            )
            text = pytesseract.image_to_string(image)

        if text.strip():
            documents.append(
                {
                    "text": text,
                    "metadata": {
                        "source": file_path.name,
                        "page": page_number,
                        "file_type": "pdf",
                    },
                }
            )

    return documents


def load_docx(file_path: Path):
    """Extract paragraphs and table rows from a Word document.

    Table cells in each row are joined using `` | `` so that labels and values
    remain associated in the text that will later be embedded.
    """
    document = Document(file_path)

    text_parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            text_parts.append(row_text)

    text = "\n".join(text_parts)

    return [
        {
            "text": text,
            "metadata": {
                "source": file_path.name,
                "file_type": "docx",
            },
        }
    ]


def load_image(file_path: Path):
    """Run Tesseract OCR on an image and return one normalized document record."""
    image = Image.open(file_path)

    text = pytesseract.image_to_string(image)

    return [
        {
            "text": text,
            "metadata": {
                "source": file_path.name,
                "file_type": "image",
            },
        }
    ]


def load_csv_or_excel(file_path: Path):
    """Convert a CSV or Excel table into embedding-friendly labeled rows.

    Each output line has the form ``column: value | ...``, retaining column
    names so retrieved table values have meaningful context.
    """
    if file_path.suffix.lower() == ".csv":
        dataframe = pd.read_csv(file_path)
    else:
        dataframe = pd.read_excel(file_path)

    text_rows = []

    for index, row in dataframe.iterrows():
        row_text = " | ".join(
            f"{column}: {row[column]}"
            for column in dataframe.columns
        )

        text_rows.append(row_text)

    return [
        {
            "text": "\n".join(text_rows),
            "metadata": {
                "source": file_path.name,
                "file_type": file_path.suffix.lower(),
            },
        }
    ]


def load_document(file_path):
    """Dispatch a source file to the appropriate format-specific loader.

    Supported extensions are TXT, PDF, DOCX, PNG/JPG/JPEG, CSV, XLSX, and XLS.
    The return value always follows the common ``text``/``metadata`` contract.

    Raises:
        ValueError: If the file extension is not supported.
    """
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension == ".txt":
        return load_txt(file_path)

    if extension == ".pdf":
        return load_pdf(file_path)

    if extension == ".docx":
        return load_docx(file_path)

    if extension in [".png", ".jpg", ".jpeg"]:
        return load_image(file_path)

    if extension in [".csv", ".xlsx", ".xls"]:
        return load_csv_or_excel(file_path)

    raise ValueError(f"Unsupported file type: {extension}")
